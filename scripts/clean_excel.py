#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""删除 Excel 中项目不存在的功能模块测试用例"""

import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

EXCEL_DIR = r"C:\Users\19795\AppData\Roaming\Tencent\Marvis\User\oAN1i2ZvLWzdAht58j6JzB69BcFc\workspace\conv_3581f9fcd3ec4d7bb6f048552949d51f\output"
xlsx_files = [f for f in os.listdir(EXCEL_DIR) if f.endswith('.xlsx')]
xlsx_path = os.path.join(EXCEL_DIR, xlsx_files[0])

wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

# 不存在的功能模块（后端无对应控制器）
REMOVE_MODULES = ['应急调度模块', '蚊媒管控模块', '车辆轨迹模块']

# 从后往前删，避免索引偏移
rows_to_delete = []
for row_idx in range(ws.max_row, 1, -1):
    module_val = ws.cell(row=row_idx, column=2).value
    if module_val and any(m in str(module_val) for m in REMOVE_MODULES):
        rows_to_delete.append(row_idx)

print(f"Removing {len(rows_to_delete)} rows from modules: {REMOVE_MODULES}")
for row_idx in rows_to_delete:
    print(f"  Deleting row {row_idx}: {ws.cell(row=row_idx, column=3).value} - {ws.cell(row=row_idx, column=4).value}")
    ws.delete_rows(row_idx)

# 重新编号
for i, row_idx in enumerate(range(2, ws.max_row + 1)):
    ws.cell(row=row_idx, column=1).value = i + 1

wb.save(xlsx_path)
print(f"\nDone! Remaining rows: {ws.max_row - 1}")
